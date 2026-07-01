"""
Generate job_description.docx for the retrieval stage
"""

from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

def create_job_description():
    doc = Document()
    
    # Title
    title = doc.add_heading('Senior AI Engineer - Search & Retrieval', 0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    
    # Company and basics
    doc.add_heading('Company', level=2)
    doc.add_paragraph('Redrob Inc.')
    
    doc.add_heading('Location', level=2)
    doc.add_paragraph('Bangalore / Hyderabad / Pune / Remote')
    
    doc.add_heading('Employment Type', level=2)
    doc.add_paragraph('Full-time')
    
    # Experience Required
    doc.add_heading('Experience Required', level=2)
    doc.add_paragraph('5-9 years of professional experience in machine learning and search systems')
    
    # What You'd Actually Be Doing
    doc.add_heading("What You'd Actually Be Doing", level=2)
    responsibilities = doc.add_paragraph(style='List Bullet')
    responsibilities.add_run('Design and optimize semantic search systems using FAISS, vector databases, and dense retrieval techniques\n')
    responsibilities = doc.add_paragraph(style='List Bullet')
    responsibilities.add_run('Build production-grade retrieval pipelines with embeddings (sentence-transformers, E5 models) and BM25 hybrid search\n')
    responsibilities = doc.add_paragraph(style='List Bullet')
    responsibilities.add_run('Implement ranking systems using learning-to-rank algorithms (LambdaRank, ListNet) and gradient boosting (LightGBM, XGBoost)\n')
    responsibilities = doc.add_paragraph(style='List Bullet')
    responsibilities.add_run('Develop RAG (Retrieval Augmented Generation) systems with LLMs using fine-tuning, LoRA, and PEFT techniques\n')
    responsibilities = doc.add_paragraph(style='List Bullet')
    responsibilities.add_run('Work with NLP models like BERT, GPT, T5, Llama, and Mistral for specialized information retrieval tasks\n')
    responsibilities = doc.add_paragraph(style='List Bullet')
    responsibilities.add_run('Deploy ML models to production using MLflow, Weights & Biases, Kubeflow, or Ray for model serving and monitoring\n')
    responsibilities = doc.add_paragraph(style='List Bullet')
    responsibilities.add_run('Optimize embeddings and search indices for performance, cost, and latency in production systems\n')
    
    # The Skills Inventory
    doc.add_heading('The Skills Inventory', level=2)
    doc.add_paragraph('Must-have Core Competencies:')
    skills = doc.add_paragraph(style='List Bullet')
    skills.add_run('FAISS, vector search, vector databases (Pinecone, Weaviate, Qdrant, Milvus)\n')
    skills = doc.add_paragraph(style='List Bullet')
    skills.add_run('Dense retrieval with sentence-transformers and E5 embeddings\n')
    skills = doc.add_paragraph(style='List Bullet')
    skills.add_run('Ranking and information retrieval (learning-to-rank, NDCG, BM25)\n')
    skills = doc.add_paragraph(style='List Bullet')
    skills.add_run('LLM and NLP production systems (transformers, fine-tuning, RAG)\n')
    skills = doc.add_paragraph(style='List Bullet')
    skills.add_run('PyTorch or TensorFlow for model development and experimentation\n')
    skills = doc.add_paragraph(style='List Bullet')
    skills.add_run('MLOps tools (MLflow, Weights & Biases, Kubeflow, Ray)\n')
    
    # Things We'd Like You To Have
    doc.add_heading('Things We\'d Like You To Have', level=2)
    nice_to_have = doc.add_paragraph(style='List Bullet')
    nice_to_have.add_run('Experience with production ML systems at scale (100k+ candidates, millions of queries)\n')
    nice_to_have = doc.add_paragraph(style='List Bullet')
    nice_to_have.add_run('Expertise in hybrid search combining dense and sparse retrieval\n')
    nice_to_have = doc.add_paragraph(style='List Bullet')
    nice_to_have.add_run('Advanced proficiency with machine learning frameworks and model optimization\n')
    nice_to_have = doc.add_paragraph(style='List Bullet')
    nice_to_have.add_run('Experience with Elasticsearch, OpenSearch, or similar search infrastructure\n')
    nice_to_have = doc.add_paragraph(style='List Bullet')
    nice_to_have.add_run('Knowledge of ONNX, TensorRT for model serving and inference optimization\n')
    
    # Things We Explicitly Do Not Want
    doc.add_heading('Things We Explicitly Do NOT Want', level=2)
    unwanted = doc.add_paragraph(style='List Bullet')
    unwanted.add_run('Computer vision specialists or image classification experts\n')
    unwanted = doc.add_paragraph(style='List Bullet')
    unwanted.add_run('Speech recognition or audio processing specialists\n')
    unwanted = doc.add_paragraph(style='List Bullet')
    unwanted.add_run('Robotics or autonomous systems background\n')
    unwanted = doc.add_paragraph(style='List Bullet')
    unwanted.add_run('Candidates with only consulting or services-company experience (TCS, Infosys, Wipro, etc.)\n')
    
    # Disqualifiers
    doc.add_heading('Disqualifiers', level=2)
    disq = doc.add_paragraph(style='List Bullet')
    disq.add_run('No demonstrated experience with embeddings, retrieval, or search systems\n')
    disq = doc.add_paragraph(style='List Bullet')
    disq.add_run('Pure consultancy background without product ML experience\n')
    disq = doc.add_paragraph(style='List Bullet')
    disq.add_run('No hands-on coding experience with Python, PyTorch, or TensorFlow\n')
    
    # Ideal Candidate Profile
    doc.add_heading('Ideal Candidate Profile', level=2)
    ideal = doc.add_paragraph()
    ideal.add_run('You are a machine learning engineer who has built search and retrieval systems at scale. You are comfortable with semantic search, embeddings, ranking algorithms, and LLM-based applications. You have strong fundamentals in machine learning, hands-on coding skills, and experience shipping production systems. You understand the tradeoffs between retrieval latency, relevance, and cost. You can optimize embedding models, manage vector indices, and build end-to-end retrieval pipelines. You care about writing clean, maintainable code and collaborating with teams on complex ML projects.')
    
    # On Location
    doc.add_heading('On Location', level=2)
    location = doc.add_paragraph()
    location.add_run('Primary focus: India-based candidates from tier-1 cities (Bangalore, Hyderabad, Pune, Mumbai, Delhi NCR, Chennai). Remote options available.')
    
    # How To Read Between The Lines
    doc.add_heading('How to Read Between the Lines', level=2)
    lines = doc.add_paragraph()
    lines.add_run('We are looking for depth in ML system design, not breadth. We want someone who understands the theory behind embeddings and ranking, and can implement production systems. Product experience is highly valued. Startup or fast-growth company backgrounds are a strong signal.')
    
    doc.save('job_description.docx')
    print("✓ Created job_description.docx")

if __name__ == '__main__':
    create_job_description()
